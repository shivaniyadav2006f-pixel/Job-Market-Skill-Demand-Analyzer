import re
import os
import uuid
import streamlit as st
import pandas as pd

from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from data_loader import find_column
from database import save_resume_record, log_activity

# =========================================================
# COMMON SKILLS
# =========================================================

COMMON_SKILLS = [
    # Programming
    "Python",
    "Java",
    "C",
    "C++",
    "C#",
    "JavaScript",
    "TypeScript",
    "R",
    "PHP",
    "Ruby",
    "Go",
    "Kotlin",
    "Swift",

    # Data
    "SQL",
    "Excel",
    "Power BI",
    "Tableau",
    "Pandas",
    "NumPy",
    "Matplotlib",
    "Seaborn",
    "Statistics",
    "Data Analysis",
    "Data Visualization",
    "Data Cleaning",
    "Data Science",

    # Machine Learning / AI
    "Machine Learning",
    "Deep Learning",
    "Artificial Intelligence",
    "TensorFlow",
    "PyTorch",
    "Scikit-learn",
    "NLP",
    "Natural Language Processing",
    "Generative AI",
    "Computer Vision",

    # Web
    "HTML",
    "CSS",
    "React",
    "Angular",
    "Vue",
    "Django",
    "Flask",
    "Node.js",
    "REST API",

    # Database
    "MySQL",
    "PostgreSQL",
    "MongoDB",
    "Oracle",
    "SQL Server",

    # Cloud
    "AWS",
    "Azure",
    "Google Cloud",
    "GCP",
    "Docker",
    "Kubernetes",
    "Linux",

    # Data Engineering
    "ETL",
    "Apache Spark",
    "Spark",
    "Hadoop",
    "Data Engineering",
    "Data Warehouse",
    "Airflow",

    # Other
    "Git",
    "GitHub",
    "GitLab",
    "PowerPoint",
    "Communication",
    "Problem Solving",
    "Agile",
    "Scrum",
]


# =========================================================
# TEXT CLEANING
# =========================================================

def clean_text(text):
    if not text:
        return ""

    text = str(text)

    text = text.replace("\n", " ")
    text = text.replace("\r", " ")
    text = re.sub(r"\s+", " ", text)

    return text.strip()


# =========================================================
# PDF READER
# =========================================================

def read_pdf(uploaded_file):

    try:
        reader = PdfReader(uploaded_file)

        pages = []

        for page in reader.pages:
            page_text = page.extract_text()

            if page_text:
                pages.append(page_text)

        return clean_text(" ".join(pages))

    except Exception as e:
        st.error(f"Unable to read PDF: {e}")
        return ""


# =========================================================
# DOCX READER
# =========================================================

def read_docx(uploaded_file):

    try:
        document = Document(uploaded_file)

        paragraphs = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        # Also read tables from DOCX resumes
        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    if cell.text.strip():
                        paragraphs.append(cell.text)

        return clean_text(" ".join(paragraphs))

    except Exception as e:
        st.error(f"Unable to read DOCX: {e}")
        return ""


# =========================================================
# RESUME TEXT EXTRACTION
# =========================================================

def extract_resume_text(uploaded_file):

    if uploaded_file is None:
        return ""

    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):

        return read_pdf(uploaded_file)

    elif file_name.endswith(".docx"):

        return read_docx(uploaded_file)

    else:

        st.error("Please upload a PDF or DOCX resume.")

        return ""

# =========================================================
# SAVE RESUME FILE
# =========================================================

def save_uploaded_resume(uploaded_file, username):

    if uploaded_file is None:
        return None

    try:

        # Create resume storage folder
        resume_folder = "uploaded_resumes"

        os.makedirs(
            resume_folder,
            exist_ok=True
        )

        # Create unique filename
        file_extension = os.path.splitext(
            uploaded_file.name
        )[1].lower()

        unique_filename = (
            f"{username}_"
            f"{uuid.uuid4().hex}"
            f"{file_extension}"
        )

        file_path = os.path.join(
            resume_folder,
            unique_filename
        )

        # Save actual file
        with open(file_path, "wb") as file:

            file.write(
                uploaded_file.getbuffer()
            )

        # Save information in database
        save_resume_record(
            username=username,
            original_filename=uploaded_file.name,
            stored_filename=file_path
        )

        # Activity log
        log_activity(
            username,
            f"Resume uploaded: {uploaded_file.name}"
        )

        return file_path

    except Exception as e:

        st.error(
            f"Unable to save resume: {e}"
        )

        return None


# =========================================================
# SKILL EXTRACTION
# =========================================================

def extract_skills(resume_text):

    text = clean_text(resume_text).lower()

    found_skills = []

    for skill in COMMON_SKILLS:

        skill_lower = skill.lower()

        # Escape special characters
        pattern = r"(?<![a-zA-Z0-9])" + re.escape(skill_lower) + r"(?![a-zA-Z0-9])"

        if re.search(pattern, text):

            found_skills.append(skill)

    return sorted(
        set(found_skills),
        key=lambda x: x.lower()
    )


# =========================================================
# JOB SKILL EXTRACTION
# =========================================================

def extract_job_skills(job_skill_text):

    if pd.isna(job_skill_text):
        return []

    text = str(job_skill_text).lower()

    found_skills = []

    for skill in COMMON_SKILLS:

        skill_lower = skill.lower()

        pattern = r"(?<![a-zA-Z0-9])" + re.escape(skill_lower) + r"(?![a-zA-Z0-9])"

        if re.search(pattern, text):

            found_skills.append(skill)

    # Also support comma-separated skills that may not
    # be present in the predefined skill list
    raw_skills = re.split(r"[,|;/]", str(job_skill_text))

    for skill in raw_skills:

        skill = skill.strip()

        if skill:

            if len(skill) <= 50:
                found_skills.append(skill)

    return sorted(
        set(found_skills),
        key=lambda x: x.lower()
    )


# =========================================================
# SKILL MATCH SCORE
# =========================================================

def calculate_skill_match(resume_skills, job_skills):

    if not job_skills:

        return 0

    resume_normalized = {
        skill.strip().lower()
        for skill in resume_skills
    }

    job_normalized = {
        skill.strip().lower()
        for skill in job_skills
    }

    matched = resume_normalized.intersection(job_normalized)

    score = (len(matched) / len(job_normalized)) * 100

    return round(score, 2)


# =========================================================
# TEXT SIMILARITY
# =========================================================

def calculate_text_similarity(resume_text, job_text):

    if not resume_text or not job_text:
        return 0

    try:

        documents = [
            resume_text,
            job_text
        ]

        vectorizer = TfidfVectorizer(
            stop_words="english"
        )

        matrix = vectorizer.fit_transform(documents)

        similarity = cosine_similarity(
            matrix[0:1],
            matrix[1:2]
        )[0][0]

        return round(similarity * 100, 2)

    except Exception:
        return 0


# =========================================================
# JOB MATCHING
# =========================================================

def match_jobs(
    df,
    resume_text,
    resume_skills,
    top_n=10
):

    if df is None or df.empty:

        return pd.DataFrame()

    # -----------------------------------------------------
    # FIND DATASET COLUMNS
    # -----------------------------------------------------

    role_column = find_column(
        df,
        [
            "Job Role",
            "Job Title",
            "Role",
            "Title",
            "Position",
            "Job"
        ]
    )

    skill_column = find_column(
        df,
        [
            "Skills",
            "Skill",
            "Required Skills",
            "Technical Skills"
        ]
    )

    company_column = find_column(
        df,
        [
            "Company",
            "Employer",
            "Organization"
        ]
    )

    location_column = find_column(
        df,
        [
            "Location",
            "City",
            "Job Location",
            "Work Location"
        ]
    )

    salary_column = find_column(
        df,
        [
            "Salary",
            "Salary Range",
            "Package",
            "CTC"
        ]
    )

    experience_column = find_column(
        df,
        [
            "Experience",
            "Experience Level",
            "Years of Experience",
            "Experience Required"
        ]
    )

    if not skill_column and not role_column:

        st.warning(
            "The dataset does not contain a Job Role or Skills column."
        )

        return pd.DataFrame()

    results = []

    # -----------------------------------------------------
    # PROCESS EACH JOB
    # -----------------------------------------------------

    for _, row in df.iterrows():

        # Job role
        job_role = ""

        if role_column:
            job_role = str(row.get(role_column, ""))

        # Job skills
        job_skills = []

        if skill_column:
            job_skills = extract_job_skills(
                row.get(skill_column, "")
            )

        # -------------------------------------------------
        # CREATE JOB TEXT
        # -------------------------------------------------

        job_parts = []

        if job_role:
            job_parts.append(job_role)

        if skill_column:
            job_parts.append(
                str(row.get(skill_column, ""))
            )

        if experience_column:
            job_parts.append(
                str(row.get(experience_column, ""))
            )

        job_text = clean_text(
            " ".join(job_parts)
        )

        # -------------------------------------------------
        # CALCULATE SCORES
        # -------------------------------------------------

        skill_score = calculate_skill_match(
            resume_skills,
            job_skills
        )

        text_score = calculate_text_similarity(
            resume_text,
            job_text
        )

        # -------------------------------------------------
        # FINAL MATCH SCORE
        # -------------------------------------------------
        #
        # Skills = 70%
        # Text similarity = 30%
        #

        if resume_skills:

            final_score = (
                (skill_score * 0.70)
                +
                (text_score * 0.30)
            )

        else:

            final_score = text_score

        final_score = round(
            min(final_score, 100),
            2
        )

        # -------------------------------------------------
        # MATCHED SKILLS
        # -------------------------------------------------

        resume_normalized = {
            skill.lower()
            for skill in resume_skills
        }

        matched_skills = []

        for skill in job_skills:

            if skill.lower() in resume_normalized:

                matched_skills.append(skill)

        # -------------------------------------------------
        # MISSING SKILLS
        # -------------------------------------------------

        missing_skills = []

        for skill in job_skills:

            if skill.lower() not in resume_normalized:

                missing_skills.append(skill)

        # -------------------------------------------------
        # STORE RESULT
        # -------------------------------------------------

        result = {
            "Job Role": job_role,
            "Match Score": final_score,
            "Matched Skills": ", ".join(matched_skills),
            "Missing Skills": ", ".join(missing_skills)
        }

        if company_column:
            result["Company"] = str(
                row.get(company_column, "")
            )

        if location_column:
            result["Location"] = str(
                row.get(location_column, "")
            )

        if salary_column:
            result["Salary"] = str(
                row.get(salary_column, "")
            )

        if experience_column:
            result["Experience"] = str(
                row.get(experience_column, "")
            )

        results.append(result)

    # -----------------------------------------------------
    # CREATE DATAFRAME
    # -----------------------------------------------------

    results_df = pd.DataFrame(results)

    if results_df.empty:
        return results_df

    results_df = results_df.sort_values(
        by="Match Score",
        ascending=False
    )

    results_df = results_df.head(top_n)

    results_df = results_df.reset_index(drop=True)

    return results_df


# =========================================================
# RESUME MATCHER TAB
# =========================================================

def resume_matcher_tab(df):

    st.subheader("📄 Resume Job Matcher")

    st.write(
        "Upload your resume to find job opportunities "
        "that match your skills and experience."
    )

    # -----------------------------------------------------
    # FILE UPLOAD
    # -----------------------------------------------------

    uploaded_resume = st.file_uploader(
        "Upload Resume",
        type=["pdf", "docx"],
        help="Upload your resume in PDF or DOCX format."
    )

    if uploaded_resume is None:

        st.info(
            "Please upload a PDF or DOCX resume to start job matching."
        )

        return
    # -----------------------------------------------------
    # SAVE RESUME
    # -----------------------------------------------------

    username = st.session_state.get(
        "username",
        ""
    )

    if not username:

        st.error(
            "User session not found. Please login again."
        )

        return

    resume_key = (
        f"{username}_"
        f"{uploaded_resume.name}_"
        f"{uploaded_resume.size}"
    )

    if st.session_state.get(
        "saved_resume_key"
    ) != resume_key:

        saved_path = save_uploaded_resume(
            uploaded_resume,
            username
        )

        if saved_path:

            st.session_state[
                "saved_resume_key"
            ] = resume_key

            st.success(
                "📄 Resume saved successfully."
            )

    # -----------------------------------------------------
    # READ RESUME
    # -----------------------------------------------------

    with st.spinner("Reading your resume..."):

        resume_text = extract_resume_text(
            uploaded_resume
        )

    if not resume_text:

        st.error(
            "No readable text was found in the uploaded resume."
        )

        return

    # -----------------------------------------------------
    # RESUME INFORMATION
    # -----------------------------------------------------

    st.success(
        "Resume uploaded and read successfully."
    )

    with st.expander("📖 View Extracted Resume Text"):

        st.write(resume_text)

    # -----------------------------------------------------
    # EXTRACT SKILLS
    # -----------------------------------------------------

    resume_skills = extract_skills(
        resume_text
    )

    st.markdown("### 🧠 Skills Detected From Resume")

    if resume_skills:

        st.write(
            ", ".join(resume_skills)
        )

        st.metric(
            "Skills Detected",
            len(resume_skills)
        )

    else:

        st.warning(
            "No predefined skills were detected in the resume."
        )

    # -----------------------------------------------------
    # FIND JOBS
    # -----------------------------------------------------

    st.markdown("### 🎯 Recommended Jobs")

    with st.spinner("Matching your resume with available jobs..."):

        matched_jobs = match_jobs(
            df=df,
            resume_text=resume_text,
            resume_skills=resume_skills,
            top_n=10
        )

    if matched_jobs.empty:

        st.warning(
            "No matching jobs could be generated from the available data."
        )

        return

    # -----------------------------------------------------
    # TOP MATCH
    # -----------------------------------------------------

    top_score = matched_jobs.iloc[0]["Match Score"]

    st.metric(
        "Best Job Match",
        f"{top_score}%"
    )

    # -----------------------------------------------------
    # JOB RESULTS
    # -----------------------------------------------------

    for index, job in matched_jobs.iterrows():

        score = float(
            job["Match Score"]
        )

        st.markdown(
            f"### {index + 1}. {job['Job Role']}"
        )

        col1, col2, col3 = st.columns(3)

        with col1:

            st.metric(
                "Match",
                f"{score}%"
            )

        with col2:

            if "Company" in job:
                st.write(
                    f"🏢 **Company:** {job['Company']}"
                )

        with col3:

            if "Location" in job:
                st.write(
                    f"📍 **Location:** {job['Location']}"
                )

        if "Salary" in job:

            st.write(
                f"💰 **Salary:** {job['Salary']}"
            )

        if "Experience" in job:

            st.write(
                f"🧑‍💻 **Experience:** {job['Experience']}"
            )

        if job["Matched Skills"]:

            st.success(
                f"✅ Matching Skills: {job['Matched Skills']}"
            )

        if job["Missing Skills"]:

            st.warning(
                f"📚 Skills to Improve: {job['Missing Skills']}"
            )

        st.divider()

    # -----------------------------------------------------
    # RESULTS TABLE
    # -----------------------------------------------------

    st.markdown("### 📊 Job Match Summary")

    st.dataframe(
        matched_jobs,
        use_container_width=True,
        hide_index=True
    )